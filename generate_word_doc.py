import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()
    
    # Page Setup - Normal Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("KARIGAR CONNECT 🏺")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x2C, 0x4A, 0x3E) # Forest Green

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(12)
    run_sub = subtitle_p.add_run("Smart India Hackathon 2026 | Additional Submission Documents")
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0xC9, 0x6F, 0x53) # Terracotta

    # Metadata Box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    col_widths = [Inches(3.2), Inches(3.3)]
    meta_data = [
        [("Problem Statement ID:", " SIH26090"), ("Ministry:", " Ministry of Social Justice and Empowerment (MoSJE)")],
        [("Team / Initiative:", " VoltEdge"), ("Target Focus:", " Livelihood Inclusion for Marginalized SC/ST Artisans")]
    ]
    
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F7F6F2")
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            bold_part, reg_part = meta_data[r_idx][c_idx]
            r1 = p.add_run(bold_part)
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r2 = p.add_run(reg_part)
            r2.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ----------------------------------------------------
    # SECTION 1: FLOW DIAGRAMS
    # ----------------------------------------------------
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Technical & User Flow Diagrams")
    run_h1.font.color.rgb = RGBColor(0x2C, 0x4A, 0x3E)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Karigar Connect is designed around zero-friction interactions for rural craftsmen. Below are the sequential flows covering product listing, fair pricing, and order fulfillment."
    )

    flow_table = doc.add_table(rows=5, cols=3)
    flow_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow_widths = [Inches(1.5), Inches(2.2), Inches(2.8)]
    
    flow_headers = ["Step / Stage", "Artisan Action", "System & AI Response"]
    for i, title in enumerate(flow_headers):
        cell = flow_table.rows[0].cells[i]
        cell.width = flow_widths[i]
        set_cell_background(cell, "2C4A3E")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    flow_steps = [
        ("1. Voice Description", "Speaks about craft in local dialect (Hindi, Tamil, Bengali).", "Speech model transcribes native audio; Gemini extracts title, SEO description, craft category, and official 8-digit HSN tax code."),
        ("2. Visual Enhancement", "Takes single smartphone picture inside workspace.", "Segmentation isolates craft object; AI replaces cluttered workshop backdrop with studio lighting, rustic wood, or festive bokeh."),
        ("3. Wage-Fair Pricing", "Enters daily hours worked, raw materials, and overheads.", "System matches input against state minimum wage laws, enforces legal wage baseline, and outputs retail (1.2x) and wholesale (1.1x) prices."),
        ("4. ONDC & Postal Pool", "Taps single 'List on ONDC' button to publish.", "Item goes live on open national buyer apps (Paytm, PhonePe). Order joins local village postal pool for India Post bulk shipping.")
    ]

    for r_idx, step in enumerate(flow_steps):
        row = flow_table.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9F8F6"
        for c_idx, text in enumerate(step):
            cell = row.cells[c_idx]
            cell.width = flow_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if c_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 2: LEAN BUSINESS MODEL CANVAS
    # ----------------------------------------------------
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. Lean Business Model Canvas")
    run_h2.font.color.rgb = RGBColor(0x2C, 0x4A, 0x3E)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Structured on a public-good, social-impact foundation under the Ministry of Social Justice and Empowerment, ensuring sustainability without taxing rural artisans."
    )

    canvas_table = doc.add_table(rows=6, cols=2)
    canvas_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_widths = [Inches(3.25), Inches(3.25)]

    canvas_data = [
        ("Problem Statement", "• High digital literacy barrier prevents rural craftsmen from e-commerce.\n• Middlemen pocket 30%–40% margins, paying sub-minimum wages.\n• Unprofessional workshop photos reduce conversion rates.\n• Isolated rural parcel deliveries are expensive."),
        ("Customer Segments", "• Primary: Rural SC/ST artisans, handloom weavers, SHGs, and folk craft clusters.\n• Secondary: Urban retail shoppers on ONDC buyer apps (Paytm, Mystore).\n• Institutional: Corporate gifting and Government e-Marketplace (GeM) buyers."),
        ("Unique Value Proposition", "• Voice-first cataloging in native dialects in under 90 seconds.\n• Automatic studio photo enhancement from raw mobile photos.\n• Legal wage protection benchmarked against state minimum wages.\n• Zero-commission direct sales on ONDC + bulk India Post pooling."),
        ("Key Solutions & Features", "• AI Vernacular Cataloger (Whisper + Gemini).\n• Background Replacement Engine (SAM 2 / Diffusion).\n• State-Wage Dynamic Pricing Assistant.\n• Vikas Khata: Offline audio cashbook with local speech synthesis."),
        ("Channels & Distribution", "• Common Service Centres (CSCs) & Village Knowledge Centres.\n• District Industries Centres (DICs) & State Handloom Directorates.\n• Self-Help Group (SHG) networks and TRIFED field offices.\n• MoSJE PM-AJAY rural artisan empowerment camps."),
        ("Revenue & Sustainability", "• 100% Free for marginalized artisans (supported by MoSJE / PM-AJAY grants).\n• 0.75% micro-convenience fee on bulk institutional B2B purchases only.\n• Premium export documentation service for international craft fairs.")
    ]

    for r_idx, (title, content) in enumerate(canvas_data):
        row = canvas_table.rows[r_idx]
        cell_t = row.cells[0]
        cell_c = row.cells[1]
        
        cell_t.width = c_widths[0]
        cell_c.width = c_widths[1]
        
        set_cell_background(cell_t, "2C4A3E")
        set_cell_background(cell_c, "F7F6F2")
        
        set_cell_margins(cell_t, 100, 100, 120, 120)
        set_cell_margins(cell_c, 100, 100, 120, 120)
        
        pt = cell_t.paragraphs[0]
        pt.paragraph_format.space_before = Pt(0)
        pt.paragraph_format.space_after = Pt(0)
        rt = pt.add_run(title)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        rt.font.size = Pt(10)
        
        pc = cell_c.paragraphs[0]
        pc.paragraph_format.space_before = Pt(0)
        pc.paragraph_format.space_after = Pt(0)
        rc = pc.add_run(content)
        rc.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 3: SYSTEM ARCHITECTURE
    # ----------------------------------------------------
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. System Architecture & Tech Visuals")
    run_h3.font.color.rgb = RGBColor(0x2C, 0x4A, 0x3E)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "A modular, light-footprint architecture designed to work reliably in low-bandwidth rural environments, with full offline persistence."
    )

    arch_table = doc.add_table(rows=4, cols=2)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    a_widths = [Inches(2.2), Inches(4.3)]
    
    arch_layers = [
        ("1. Presentation Tier (Mobile & Web Client)", "Responsive HTML5/Tailwind SPA with zero heavy dependencies. Supports offline transaction caching in LocalStorage, native microphone capture, and Web Speech API (SpeechSynthesis) for vernacular voice readouts in Hindi/Tamil."),
        ("2. Application API (FastAPI Backend)", "High-performance Python FastAPI server managing REST endpoints for catalog generation, fair wage compliance checks, postal aggregation logic, and batch offline-to-cloud synchronization."),
        ("3. AI & Vision Pipeline (Inference Layer)", "• Indic Speech: Fine-tuned Whisper models transcribing regional dialects.\n• Metadata Engine: Gemini 1.5 Flash generating English titles, descriptions, and 8-digit HSN codes.\n• Vision Studio: Object segmentation and background placement for studio-quality photos."),
        ("4. Ecosystem Networks (External Gateways)", "• ONDC Protocol: Direct integration with open Beckn protocol registry for discovery across consumer apps.\n• India Post Hub: Postal pooling algorithm aggregating cluster packages by village PIN code.")
    ]

    for r_idx, (layer_title, layer_desc) in enumerate(arch_layers):
        row = arch_table.rows[r_idx]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width, c2.width = a_widths[0], a_widths[1]
        
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9F8F6"
        set_cell_background(c1, "F2EFE9")
        set_cell_background(c2, bg)
        
        set_cell_margins(c1, 100, 100, 100, 100)
        set_cell_margins(c2, 100, 100, 100, 100)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(layer_title)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x2C, 0x4A, 0x3E)
        
        p2 = c2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(layer_desc)
        r2.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 4: ROADMAP
    # ----------------------------------------------------
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. Product Implementation Roadmap")
    run_h4.font.color.rgb = RGBColor(0x2C, 0x4A, 0x3E)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "A realistic 3-phase rollout designed to progress from district craft clusters to national institutional procurement."
    )

    road_table = doc.add_table(rows=4, cols=3)
    road_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_widths = [Inches(1.4), Inches(2.2), Inches(2.9)]

    r_headers = ["Phase & Timeline", "Key Milestones", "Target Deliverables & KPIs"]
    for i, title in enumerate(r_headers):
        cell = road_table.rows[0].cells[i]
        cell.width = r_widths[i]
        set_cell_background(cell, "C96F53") # Terracotta
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    phases = [
        ("Phase 1\n(Months 1–6)\nCluster Pilot", "Deploy working prototype in 5 high-density artisan clusters (Jaipur pottery, Madhubani art, Assam cane).", "• 2,500 active artisans onboarded.\n• 10,000 product sheets generated.\n• Zero-failure offline synchronization verified on 2G/3G networks."),
        ("Phase 2\n(Months 7–18)\nONDC & Postal Scale", "Direct India Post API dispatch integration; expand vernacular voice support to 12 scheduled Indian languages.", "• 50,000 artisans across 8 states.\n• 40% reduction in average rural parcel shipping costs.\n• Bank micro-lending pilot using Vikas Khata cashflow logs."),
        ("Phase 3\n(Months 19–36)\nInstitutional Linkage", "Connect directly to the Government e-Marketplace (GeM) for public procurement; deploy on-device Edge AI.", "• Direct supply pipelines to PSUs and corporate CSR buyers.\n• Zero-data-cost on-device background removal on budget smartphones.\n• Fully integrated into MoSJE PM-AJAY welfare programs.")
    ]

    for r_idx, phase in enumerate(phases):
        row = road_table.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9F8F6"
        for c_idx, text in enumerate(phase):
            cell = row.cells[c_idx]
            cell.width = r_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if c_idx == 0:
                r.font.bold = True

    # Save to target destination
    output_path = r"c:\Users\DELL\Downloads\Karigar-Connect-main\Karigar-Connect-main\Karigar_Connect_SIH2026_Additional_Documents.docx"
    doc.save(output_path)
    print("Document successfully created at:", output_path)

if __name__ == "__main__":
    create_document()
